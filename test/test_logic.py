"""로직 테스트 — python3 -m unittest discover -s test

기대값은 예시 문서(stdreview/sample_docx.py)에 일부러 넣은 결함을 손으로 세어 적었습니다.
"""

import os
import sys
import tempfile
import unittest
from collections import Counter
from pathlib import Path

ROOT = Path(__file__).resolve().parent.parent
sys.path.insert(0, str(ROOT))

from stdreview import defaults  # noqa: E402
from stdreview.checks import (assemble, captions, find_std_numbers, j, run_checks, summarize,  # noqa: E402
                              term_entries, title_keywords)
from stdreview.config import diff, validate  # noqa: E402
from stdreview.docx_reader import build, parse_target, read_docx  # noqa: E402
from stdreview.improve import candidates  # noqa: E402
from stdreview.prompts import build_all  # noqa: E402
from stdreview.sample_docx import TITLE, bad_blocks, build_docx, good_blocks  # noqa: E402

try:
    import openpyxl  # noqa: F401
    HAS_XL = True
except ImportError:
    HAS_XL = False


def make(blocks, tmp, name="t.docx"):
    p = Path(tmp) / name
    build_docx(p, TITLE, blocks)
    return read_docx(p)


class ReaderTest(unittest.TestCase):
    def setUp(self):
        self.tmp = tempfile.mkdtemp()

    def test_headings_and_title(self):
        doc = make(good_blocks(), self.tmp)
        self.assertEqual(doc["title"], TITLE)
        nums = [h["num"] for h in doc["headings"]]
        self.assertEqual(nums, ["1", "2", "3", "4", "4.1", "4.2", "4.3", "5", "6", "7", "8"])
        self.assertEqual(sum(b.get("omath", 0) for b in doc["blocks"]), 1)
        self.assertEqual(sum(1 for b in doc["blocks"] if b.get("figure")), 1)

    def test_auto_numbered_headings(self):
        # Word 자동 번호 제목(글자에 번호가 없음)은 스타일 수준으로 번호를 셉니다.
        blocks = [
            {"kind": "p", "text": "제목", "style": "Title", "style_level": None},
            {"kind": "p", "text": "목적 및 정의", "style": "heading 1", "style_level": 1},
            {"kind": "p", "text": "본문", "style": "", "style_level": None},
            {"kind": "p", "text": "설계절차", "style": "heading 1", "style_level": 1},
            {"kind": "p", "text": "흐름", "style": "heading 2", "style_level": 2},
            {"kind": "p", "text": "계산", "style": "heading 2", "style_level": 2},
            {"kind": "p", "text": "검증", "style": "heading 1", "style_level": 1},
        ]
        doc = build(blocks)
        self.assertEqual([h["num"] for h in doc["headings"]], ["1", "2", "2.1", "2.2", "3"])
        self.assertEqual(doc["blocks"][2]["sec"], "1")

    def test_numbered_list_is_not_heading(self):
        # 2항 안의 「1. EXS-1001 …」 목록 줄은 제목이 아닙니다(번호가 앞 항보다 작음).
        blocks = [
            {"kind": "p", "text": "1. 목적", "style": "", "style_level": None},
            {"kind": "p", "text": "2. 관련 표준", "style": "", "style_level": None},
            {"kind": "p", "text": "1. EXS-1001 예시 기준", "style": "", "style_level": None},
            {"kind": "p", "text": "3. 용어정의", "style": "", "style_level": None},
            {"kind": "p", "text": "3.1 온도는 85 ℃ 로 한다.", "style": "", "style_level": None},
        ]
        doc = build(blocks)
        self.assertEqual([h["num"] for h in doc["headings"]], ["1", "2", "3"])
        self.assertEqual(doc["blocks"][2]["sec"], "2")

    def test_bad_file(self):
        from stdreview.docx_reader import DocxError
        p = Path(self.tmp) / "x.docx"
        p.write_text("not a zip")
        with self.assertRaises(DocxError):
            read_docx(p)

    def test_parse_target(self):
        _, m = parse_target("1,4+")
        self.assertTrue(m("1"))
        self.assertTrue(m("4.2"))
        self.assertTrue(m("7"))
        self.assertFalse(m("2"))
        self.assertFalse(m("3.1"))


class CheckTest(unittest.TestCase):
    def setUp(self):
        self.tmp = tempfile.mkdtemp()
        self.cfg = defaults.default_config()

    def rows(self, blocks):
        doc = make(blocks, self.tmp)
        return doc, assemble(run_checks(doc, self.cfg, "설계표준"), self.cfg)

    def test_good_doc_has_no_findings(self):
        _, rows = self.rows(good_blocks())
        bad = [r for r in rows if r["result"] in ("부적합", "권고")]
        self.assertEqual(bad, [])
        s = summarize(rows)
        self.assertEqual(s["부적합"], 0)
        self.assertEqual(s["권고"], 0)

    def test_bad_doc_findings(self):
        _, rows = self.rows(bad_blocks())
        got = Counter((r["no"], r["result"]) for r in rows if r["result"] in ("부적합", "권고"))
        expected = Counter({
            ("공통-1", "부적합"): 1,   # 6항 없음
            ("공통-1", "권고"): 1,     # 8항(선택) 없음
            ("2-1", "부적합"): 2,      # 번호 없는 줄, 제목 없는 번호
            ("3-1", "부적합"): 1,      # 방열판 설명 비어 있음
            ("3-2", "부적합"): 2,      # 이씨유, Ecu
            ("3-2", "권고"): 1,        # 냉각팬 본문 미사용
            ("4-1", "부적합"): 1,      # 4.2 빠짐
            ("4-2", "권고"): 1,        # 흐름도 캡션 낱말 없음
            ("4-3", "권고"): 1,        # 수식 없음
            ("4-4", "부적합"): 3,      # 캡션 없는 표, 그림 2 빠짐, 없는 표 5 참조
            ("4-6", "부적합"): 1,      # EXS-9001 이 2항 목록에 없음
            ("5-2", "부적합"): 1,      # 「양호할 것」
            ("6-1", "부적합"): 1,      # 6항 없음
            ("8-1", "권고"): 1,        # 8항 없음
        })
        self.assertEqual(got, expected)
        s = summarize(rows)
        self.assertEqual((s["부적합"], s["권고"]), (13, 5))

    def test_fix_message_uses_template(self):
        _, rows = self.rows(bad_blocks())
        r = next(r for r in rows if r["no"] == "3-2" and "이씨유" in r["detail"])
        self.assertEqual(r["detail"], "용어 「ECU」가 「이씨유」로 1곳에 다르게 쓰였습니다.")
        self.assertEqual(r["fix"], "Template 3항 용어정의 ↔ 본문 전체 기준에 따라 「이씨유」 표기를 용어정의대로 「ECU」로 통일해 주시기 바랍니다.")
        self.assertIn("4.1절", r["location"])
        ok = [r for r in rows if r["result"] in ("적합", "판단 필요")]
        self.assertTrue(all(r["fix"] == "" for r in ok))

    def test_inactive_criterion_is_skipped(self):
        for c in self.cfg["criteria"]:
            if c["no"] == "4-3":
                c["active"] = False
        _, rows = self.rows(bad_blocks())
        self.assertFalse(any(r["no"] == "4-3" for r in rows))

    def test_std_numbers(self):
        cfg = self.cfg
        got = [g for _, _, g in find_std_numbers("ISO 26262-1:2018 과 KS C IEC 61508-3, EXS-1001 을 따른다", cfg)]
        self.assertEqual(got, ["ISO 26262-1:2018", "KS C IEC 61508-3", "EXS-1001"])
        self.assertEqual(find_std_numbers("온도 85 ℃", cfg), [])

    def test_terms_and_captions(self):
        doc = make(good_blocks(), self.tmp)
        self.assertEqual([e["term"] for e in term_entries(doc)], ["ECU", "방열판", "접합부 온도"])
        caps, objects, _ = captions(doc, self.cfg)
        self.assertEqual([(c["kind"], c["num"]) for c in caps], [("표", "1"), ("표", "2"), ("그림", "1"), ("표", "3")])
        self.assertTrue(all(o["caption"] for o in objects))

    def test_sentence_starting_with_table_ref_is_not_caption(self):
        doc = build([
            {"kind": "p", "text": "표 1. 입력값", "style": "caption", "style_level": None},
            {"kind": "p", "text": "표 1의 입력값으로 다시 계산해 확인한다.", "style": "", "style_level": None},
        ])
        caps, _, _ = captions(doc, self.cfg)
        self.assertEqual(len(caps), 1)

    def test_josa(self):
        self.assertEqual(j("표", "이가"), "표가")
        self.assertEqual(j("그림", "이가"), "그림이")
        self.assertEqual(j("「ECU」", "으로"), "「ECU」로")
        self.assertEqual(j("「용어정의」", "으로"), "「용어정의」로")
        self.assertEqual(j("「방열판」", "으로"), "「방열판」으로")
        self.assertEqual(j("「표 3」", "이가"), "「표 3」이")

    def test_title_keywords(self):
        self.assertEqual(title_keywords(TITLE), ["전장", "제어기", "ECU", "방열"])


class PromptTest(unittest.TestCase):
    def test_prompts_for_judgment_criteria(self):
        cfg = defaults.default_config()
        tmp = tempfile.mkdtemp()
        doc = make(good_blocks(), tmp)
        rows = assemble(run_checks(doc, cfg, "설계표준"), cfg)
        prompts = dict(build_all(cfg, doc, "설계표준", rows, {"1-1": ["(예) 목적이 막연함"]}))
        self.assertEqual(sorted(prompts), ["1-1", "3-1", "4-5", "4-6", "5-1", "5-2", "7-1"])
        p = prompts["4-5"]
        self.assertIn("1 목적 및 정의", p)
        self.assertIn("4.2 접합부 온도 계산", p)
        self.assertNotIn("관련 표준 및 인용", p)      # 2항은 대상이 아님
        self.assertIn("판정: 적합 | 부적합 | 권고", p)
        self.assertIn("(예) 목적이 막연함", prompts["1-1"])

    def test_user_added_criterion_gets_prompt(self):
        cfg = defaults.default_config()
        cfg["criteria"].append({"no": "9-1", "item": "7항", "text": "(예시) 보호구 기준을 적을 것", "method": "판단(프롬프트)",
                                "code": "", "target": "7", "active": True, "fix": "", "prompt": "", "note": ""})
        tmp = tempfile.mkdtemp()
        doc = make(good_blocks(), tmp)
        rows = assemble(run_checks(doc, cfg, "설계표준"), cfg)
        self.assertTrue(any(r["no"] == "9-1" and r["result"] == "판단 필요" for r in rows))
        prompts = dict(build_all(cfg, doc, "설계표준", rows))
        self.assertIn("모따기", prompts["9-1"])


class ConfigTest(unittest.TestCase):
    def test_default_is_valid(self):
        self.assertEqual(validate(defaults.default_config()), [])

    def test_validate_catches_errors(self):
        cfg = defaults.default_config()
        cfg["criteria"][2]["code"] = "NO_SUCH"
        cfg["criteria"][3]["fix"] = "고쳐 주세요"
        cfg["std_patterns"].append({"name": "x", "regex": "("})
        cfg["columns"] = [c for c in cfg["columns"] if c["key"] != "verdict"]
        errs = validate(cfg)
        self.assertEqual(len(errs), 4)

    def test_diff(self):
        old = defaults.default_config()
        new = defaults.default_config()
        new["criteria"][4]["text"] = "고친 문구"
        new["criteria"][8]["active"] = False
        new["criteria"].append({"no": "9-1", "item": "x", "text": "새 기준", "method": "판단(프롬프트)", "code": "",
                                "target": "7", "active": True, "fix": "", "prompt": "", "note": ""})
        new["variants"].append({"term": "방열판", "variants": ["히트싱크"]})
        new["settings"]["보완후보_제외비율"] = "0.3"
        kinds = Counter(c["kind"] for c in diff(old, new))
        self.assertEqual(kinds, Counter({"수정": 2, "사용 중지": 1, "추가": 2}))

    @unittest.skipUnless(HAS_XL, "openpyxl 없음")
    def test_xlsx_roundtrip(self):
        from stdreview.config import read_xlsx, write_xlsx
        cfg = defaults.default_config()
        p = Path(tempfile.mkdtemp()) / "기준표.xlsx"
        write_xlsx(cfg, p)
        back = read_xlsx(p)
        self.assertEqual(diff(cfg, back), [])
        self.assertEqual(validate(back), [])


class ImproveTest(unittest.TestCase):
    def test_candidates(self):
        fb = (
            [{"no": "4-3", "verdict": "제외", "comment": ""}] * 3
            + [{"no": "4-3", "verdict": "채택", "comment": ""}]           # 4-3: 제외 3/4 = 0.75 → 후보
            + [{"no": "4-4", "verdict": "제외", "comment": ""}]
            + [{"no": "4-4", "verdict": "채택", "comment": ""}] * 3        # 4-4: 제외 1/4 = 0.25 → 아님
            + [{"no": "3-2", "verdict": "채택", "comment": "EUC 도 사전에 추가"}]  # 의견 → 후보
            + [{"no": "2-1", "verdict": "제외", "comment": ""}] * 2         # 2건 < 최소 3건 → 아님
        )
        stats = candidates(fb, defaults.CRITERIA, min_count=3, ratio=0.5)
        by = {s["no"]: s for s in stats}
        self.assertTrue(by["4-3"]["candidate"])
        self.assertAlmostEqual(by["4-3"]["ratio"], 0.75)
        self.assertFalse(by["4-4"]["candidate"])
        self.assertTrue(by["3-2"]["candidate"])
        self.assertFalse(by["2-1"]["candidate"])
        self.assertEqual([s["no"] for s in stats if s["candidate"]], ["3-2", "4-3"])


@unittest.skipUnless(HAS_XL, "openpyxl 없음")
class FlowTest(unittest.TestCase):
    """점검 → 회의록 초안 → 판정 적기 → 되받기 → 보완 후보 → 기준표 수정까지 한 바퀴."""

    def test_full_cycle(self):
        import openpyxl
        import review_tool
        from stdreview.store import Store
        from stdreview.minutes import read_feedback
        tmp = Path(tempfile.mkdtemp())
        src = tmp / "가상표준.docx"
        build_docx(src, TITLE, bad_blocks())
        store = Store(tmp / "t.db")
        res = review_tool.do_review(store, src, "설계표준", tmp / "out", quiet=True)
        # 검토자가 판정 열 이름을 「판정」으로 바꾸고 4-3 을 제외했다고 가정
        wb = openpyxl.load_workbook(res["xlsx"])
        ws = wb["회의록 초안"]
        head = [c.value for c in ws[1]]
        ci = {h: i for i, h in enumerate(head)}
        ws.cell(1, ci["검토자 판정"] + 1).value = "판정"
        for row in ws.iter_rows(min_row=2):
            if row[ci["결과"]].value in ("부적합", "권고"):
                row[ci["검토자 판정"]].value = "제외" if row[ci["No"]].value == "4-3" else "채택"
                if row[ci["No"]].value == "4-3":
                    row[ci["보완 의견"]].value = "수식이 필요 없는 절차도 있음"
        wb.save(res["xlsx"])
        _, cfg = store.current()
        with self.assertRaises(Exception):
            read_feedback(res["xlsx"], cfg)                      # 열 이름이 바뀌어 못 찾음
        rid, items, warn = read_feedback(res["xlsx"], cfg, {"verdict": "판정"})
        self.assertEqual(rid, res["id"])
        store.upsert_feedback(rid, items)
        store.upsert_feedback(rid, items)                         # 두 번 넣어도 중복 안 됨
        fb = store.all_feedback()
        self.assertEqual(len(fb), len(items))
        self.assertEqual(sum(1 for f in fb if f["verdict"] == "채택"), 17)
        stats = candidates(fb, cfg["criteria"], 1, 0.5)
        self.assertEqual([s["no"] for s in stats if s["candidate"]], ["4-3"])
        # 기준표 보완: 4-3 사용 중지 → 새 버전, 다음 점검에서 4-3 행이 없어짐
        import copy
        new = copy.deepcopy(cfg)
        for c in new["criteria"]:
            if c["no"] == "4-3":
                c["active"] = False
        vid = review_tool.apply_new_config(store, new, "4-3 권고 제외", "테스트", assume_yes=True)
        self.assertEqual(vid, 1)
        self.assertEqual(store.history()[0]["changes"][0]["kind"], "사용 중지")
        res2 = review_tool.do_review(store, src, "설계표준", tmp / "out2", quiet=True)
        self.assertFalse(any(r["no"] == "4-3" for r in res2["rows"]))
        # 채택된 지적이 다음 판단 프롬프트의 과거 사례로 들어감(3-1 방열판 설명 누락)
        p31 = dict(res2["prompts"])["3-1"]
        self.assertIn("방열판", p31.split("[과거 심의에서 채택된 지적 예")[1])
        store.close()


if __name__ == "__main__":
    unittest.main()
