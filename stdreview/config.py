"""기준표 — 검사·비교(순수 함수)와 엑셀 읽기·쓰기.

기준표 엑셀 시트: 기준표 / 항목 / 이형표기 / 정성표현 / 표준번호형식 / 설정 / 회의록열 / 안내
검토자가 이 엑셀을 고쳐 다시 넣으면 새 버전이 됩니다(기획서 5.3 평가 기준 보완).
"""

import re

from . import defaults
from .checks import KNOWN_CODES


class ConfigError(Exception):
    pass


def _yn(v):
    s = str(v if v is not None else "").strip().upper()
    return s in ("Y", "YES", "예", "O", "TRUE", "1", "사용")


def validate(cfg):
    """문제 목록을 돌려줍니다(비었으면 정상)."""
    errs = []
    nos = [c.get("no") for c in cfg.get("criteria", [])]
    for n in sorted({n for n in nos if nos.count(n) > 1}):
        errs.append("기준표: No 「%s」가 두 번 이상 있습니다." % n)
    for c in cfg.get("criteria", []):
        if not c.get("no"):
            errs.append("기준표: No 가 빈 행이 있습니다.")
        code = (c.get("code") or "").strip()
        if code and code not in KNOWN_CODES:
            errs.append("기준표 %s: 점검코드 「%s」는 도구에 없는 코드입니다. 비워 두면 판단 프롬프트 기준으로 씁니다." % (c.get("no"), code))
        if c.get("active") and not code and "판단" not in str(c.get("method", "")):
            errs.append("기준표 %s: 점검코드가 없으면 점검 방식에 「판단」이 들어가야 합니다(판단 프롬프트로 확인)." % c.get("no"))
        fix = c.get("fix") or ""
        if fix and "{내용}" not in fix:
            errs.append("기준표 %s: 수정 요청 문구에 {내용} 자리가 없습니다." % c.get("no"))
    for p in cfg.get("std_patterns", []):
        try:
            re.compile(p.get("regex", ""))
        except re.error as e:
            errs.append("표준번호형식 「%s」 정규식 오류: %s" % (p.get("name"), e))
    try:
        re.compile(str(cfg.get("settings", {}).get("수치_단위_정규식", "")))
    except re.error as e:
        errs.append("설정 수치_단위_정규식 오류: %s" % e)
    for key in ("보완후보_최소판정건수", "보완후보_제외비율", "짧은_설명_글자수", "프롬프트_본문_최대글자", "과거사례_최대건수"):
        v = cfg.get("settings", {}).get(key)
        if v is not None and str(v).strip() != "":
            try:
                float(v)
            except ValueError:
                errs.append("설정 %s 값 「%s」가 숫자가 아닙니다." % (key, v))
    keys = [c.get("key") for c in cfg.get("columns", [])]
    for k in ("seq", "no", "result", "verdict", "comment"):
        if k not in keys:
            errs.append("회의록열: 키 「%s」가 없습니다(필수)." % k)
    heads = [c.get("header") for c in cfg.get("columns", [])]
    for h in sorted({h for h in heads if heads.count(h) > 1}):
        errs.append("회의록열: 열 이름 「%s」가 겹칩니다." % h)
    for it in cfg.get("items", []):
        if it.get("type") not in defaults.STD_TYPES:
            errs.append("항목: 표준 종류 「%s」는 %s 중 하나여야 합니다." % (it.get("type"), ", ".join(defaults.STD_TYPES)))
    return errs


CRIT_FIELDS = [("item", "Template 항목"), ("text", "평가 기준"), ("method", "점검 방식"), ("code", "점검코드"),
               ("target", "대상 절"), ("active", "사용"), ("fix", "수정 요청 문구"), ("prompt", "판단 프롬프트 지시"), ("note", "비고")]


def diff(old, new):
    """두 기준표의 차이. [{kind, no, field, before, after}] — kind: 추가/삭제/수정/사용 중지/다시 사용."""
    out = []
    o = {c["no"]: c for c in old.get("criteria", [])}
    n = {c["no"]: c for c in new.get("criteria", [])}
    for no, c in n.items():
        if no not in o:
            out.append({"kind": "추가", "no": no, "field": "기준", "before": "", "after": c.get("text", "")})
            continue
        for f, label in CRIT_FIELDS:
            a, b = o[no].get(f), c.get(f)
            if f == "active":
                if bool(a) != bool(b):
                    out.append({"kind": "사용 중지" if not b else "다시 사용", "no": no, "field": label,
                                "before": "Y" if a else "N", "after": "Y" if b else "N"})
            elif str(a or "") != str(b or ""):
                out.append({"kind": "수정", "no": no, "field": label, "before": str(a or ""), "after": str(b or "")})
    for no in o:
        if no not in n:
            out.append({"kind": "삭제", "no": no, "field": "기준", "before": o[no].get("text", ""), "after": ""})

    def rows_of(cfg, key, fmt):
        return [fmt(r) for r in cfg.get(key, [])]

    sheets = [
        ("items", "항목", lambda r: "%s %s항 %s 필수=%s" % (r.get("type"), r.get("num"), r.get("name"), "Y" if r.get("required") else "N")),
        ("variants", "이형표기", lambda r: "%s ← %s" % (r.get("term"), ", ".join(r.get("variants", [])))),
        ("qualitative", "정성표현", lambda r: str(r)),
        ("std_patterns", "표준번호형식", lambda r: "%s: %s" % (r.get("name"), r.get("regex"))),
        ("columns", "회의록열", lambda r: "%s=%s" % (r.get("key"), r.get("header"))),
    ]
    for key, label, fmt in sheets:
        a, b = rows_of(old, key, fmt), rows_of(new, key, fmt)
        for x in b:
            if x not in a:
                out.append({"kind": "추가", "no": label, "field": label, "before": "", "after": x})
        for x in a:
            if x not in b:
                out.append({"kind": "삭제", "no": label, "field": label, "before": x, "after": ""})
    so, sn = old.get("settings", {}), new.get("settings", {})
    for k in sorted(set(so) | set(sn)):
        if str(so.get(k, "")) != str(sn.get(k, "")):
            out.append({"kind": "수정" if k in so and k in sn else ("추가" if k in sn else "삭제"),
                        "no": "설정", "field": k, "before": str(so.get(k, "")), "after": str(sn.get(k, ""))})
    return out


def items_from_template(doc, std_type):
    """Template .docx 의 상위 제목을 그 종류의 항목 목록으로 만듭니다."""
    tops = [h for h in doc["headings"] if h["level"] == 1]
    return [{"type": std_type, "order": i, "num": h["num"], "name": h["title"], "required": True}
            for i, h in enumerate(tops, 1)]


# ---------- 엑셀 ----------

def _openpyxl():
    try:
        import openpyxl  # noqa: F401
        return openpyxl
    except ImportError:
        raise ConfigError("openpyxl 이 설치되어 있지 않습니다. README 「실행 방법」의 폐쇄망 설치 안내를 따라 주세요.")


CRIT_HEAD = ["No", "Template 항목", "평가 기준", "점검 방식", "점검코드", "대상 절", "사용(Y/N)", "수정 요청 문구", "판단 프롬프트 지시", "비고"]
ITEM_HEAD = ["표준 종류", "순서", "항목 번호", "항목명(| 로 여러 이름)", "필수(Y/N)"]


def _style_sheet(ws, widths):
    from openpyxl.styles import Alignment, Font, PatternFill
    for i, w in enumerate(widths, 1):
        ws.column_dimensions[ws.cell(1, i).column_letter].width = w
    for c in ws[1]:
        c.font = Font(bold=True, color="FFFFFF")
        c.fill = PatternFill("solid", fgColor="1F4E79")
        c.alignment = Alignment(vertical="center", wrap_text=True)
    for row in ws.iter_rows(min_row=2):
        for c in row:
            c.alignment = Alignment(vertical="top", wrap_text=True)
    ws.freeze_panes = "A2"


def write_xlsx(cfg, path, version_label=""):
    openpyxl = _openpyxl()
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "안내"
    guide = [
        ["기술표준 검토 도구 — 기준표"],
        ["버전", version_label],
        [""],
        ["고치는 법"],
        ["1. 「기준표」 시트에서 문구를 고치거나, 새 행을 추가하거나, 사용(Y/N)을 N 으로 바꿉니다."],
        ["2. 새 기준은 점검코드를 비우고 점검 방식에 「판단(프롬프트)」를 적으면 판단 프롬프트가 만들어집니다. 대상 절에 「1」「4+」(4항 이후 전체)「1,4+」처럼 적습니다."],
        ["3. 수정 요청 문구의 {항목} 에는 Template 항목, {내용} 에는 도구가 만든 고칠 내용이 들어갑니다."],
        ["4. 저장한 뒤 python review_tool.py criteria import 이파일.xlsx --reason \"고친 이유\" 로 다시 넣습니다. 바뀐 점을 먼저 보여 주고 확정합니다."],
        ["5. 점검코드는 도구에 들어 있는 규칙 이름입니다. 없는 이름을 적으면 넣을 때 알려 줍니다: " + ", ".join(sorted(KNOWN_CODES))],
    ]
    for r in guide:
        ws.append(r)
    ws.column_dimensions["A"].width = 30
    ws.column_dimensions["B"].width = 60

    ws = wb.create_sheet("기준표")
    ws.append(CRIT_HEAD)
    for c in cfg["criteria"]:
        ws.append([c.get("no"), c.get("item"), c.get("text"), c.get("method"), c.get("code"), c.get("target"),
                   "Y" if c.get("active") else "N", c.get("fix"), c.get("prompt"), c.get("note")])
    _style_sheet(ws, [9, 18, 48, 16, 18, 10, 9, 36, 48, 36])

    ws = wb.create_sheet("항목")
    ws.append(ITEM_HEAD)
    for it in cfg["items"]:
        ws.append([it.get("type"), it.get("order"), str(it.get("num")), it.get("name"), "Y" if it.get("required") else "N"])
    _style_sheet(ws, [16, 8, 10, 36, 10])

    ws = wb.create_sheet("이형표기")
    ws.append(["표준 용어", "허용하지 않는 표기(쉼표 구분)"])
    for v in cfg["variants"]:
        ws.append([v.get("term"), ", ".join(v.get("variants", []))])
    _style_sheet(ws, [20, 50])

    ws = wb.create_sheet("정성표현")
    ws.append(["정성 표현 낱말"])
    for w in cfg["qualitative"]:
        ws.append([w])
    _style_sheet(ws, [30])

    ws = wb.create_sheet("표준번호형식")
    ws.append(["이름", "정규식"])
    for p in cfg["std_patterns"]:
        ws.append([p.get("name"), p.get("regex")])
    _style_sheet(ws, [40, 60])

    ws = wb.create_sheet("설정")
    ws.append(["키", "값", "설명"])
    for k, v in cfg["settings"].items():
        ws.append([k, str(v), defaults.SETTING_NOTES.get(k, "")])
    _style_sheet(ws, [24, 50, 60])

    ws = wb.create_sheet("회의록열")
    ws.append(["키(바꾸지 마세요)", "열 이름(기존 회의록 양식에 맞춰 바꿉니다)", "설명"])
    for c in cfg["columns"]:
        ws.append([c.get("key"), c.get("header"), defaults.COLUMN_NOTES.get(c.get("key"), "")])
    _style_sheet(ws, [18, 36, 50])
    wb.save(path)


def _rows(wb, name):
    if name not in wb.sheetnames:
        raise ConfigError("기준표 엑셀에 「%s」 시트가 없습니다." % name)
    ws = wb[name]
    out = []
    for i, row in enumerate(ws.iter_rows(values_only=True)):
        if i == 0:
            continue
        vals = ["" if v is None else str(v).strip() for v in row]
        if any(vals):
            out.append(vals)
    return out


def _get(row, i):
    return row[i] if i < len(row) else ""


def read_xlsx(path):
    openpyxl = _openpyxl()
    try:
        wb = openpyxl.load_workbook(path, data_only=True)
    except Exception as e:
        raise ConfigError("기준표 엑셀을 열 수 없습니다: %s (%s)" % (path, e))
    cfg = {"criteria": [], "items": [], "variants": [], "qualitative": [], "std_patterns": [], "settings": {}, "columns": []}
    for r in _rows(wb, "기준표"):
        cfg["criteria"].append({
            "no": _get(r, 0), "item": _get(r, 1), "text": _get(r, 2), "method": _get(r, 3), "code": _get(r, 4),
            "target": _get(r, 5), "active": _yn(_get(r, 6)), "fix": _get(r, 7), "prompt": _get(r, 8), "note": _get(r, 9),
        })
    for r in _rows(wb, "항목"):
        try:
            order = int(float(_get(r, 1) or 0))
        except ValueError:
            order = 0
        num = _get(r, 2)
        if num.endswith(".0"):
            num = num[:-2]
        cfg["items"].append({"type": _get(r, 0), "order": order, "num": num, "name": _get(r, 3), "required": _yn(_get(r, 4))})
    for r in _rows(wb, "이형표기"):
        cfg["variants"].append({"term": _get(r, 0), "variants": [x.strip() for x in _get(r, 1).split(",") if x.strip()]})
    cfg["qualitative"] = [r[0] for r in _rows(wb, "정성표현") if r and r[0]]
    for r in _rows(wb, "표준번호형식"):
        cfg["std_patterns"].append({"name": _get(r, 0), "regex": _get(r, 1)})
    for r in _rows(wb, "설정"):
        cfg["settings"][_get(r, 0)] = _get(r, 1)
    for r in _rows(wb, "회의록열"):
        cfg["columns"].append({"key": _get(r, 0), "header": _get(r, 1)})
    return cfg
