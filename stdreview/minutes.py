"""회의록 초안 엑셀 쓰기와 검토자 판정 되받기."""

from .config import ConfigError, _openpyxl
from . import defaults

INFO_SHEET = "정보"
MAIN_SHEET = "회의록 초안"


def headers(cfg):
    return [(c["key"], c["header"]) for c in cfg.get("columns", []) if c.get("key") in defaults.COLUMN_KEYS]


def write_minutes(path, rows, cfg, info, prompts=None):
    openpyxl = _openpyxl()
    from openpyxl.styles import Alignment, Font, PatternFill
    from openpyxl.worksheet.datavalidation import DataValidation

    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = MAIN_SHEET
    cols = headers(cfg)
    start = 1
    if info.get("예시 데이터") == "예":
        ws.append(["예시 데이터 — 가상 표준 문서로 만든 시연용 결과입니다. 실제 심의 결과가 아닙니다."])
        ws.cell(1, 1).font = Font(bold=True, color="9C0006")
        ws.cell(1, 1).fill = PatternFill("solid", fgColor="FFC7CE")
        start = 2
    ws.append([h for _, h in cols])
    fills = {"부적합": "FFC7CE", "권고": "FFEB9C", "판단 필요": "DDEBF7", "적합": "E2EFDA"}
    for r in rows:
        ws.append([r.get(k, "") for k, _ in cols])
    widths = {"seq": 6, "no": 8, "item": 18, "criterion": 36, "result": 10, "detail": 50, "location": 30, "fix": 46, "verdict": 12, "comment": 36}
    for i, (k, _) in enumerate(cols, 1):
        ws.column_dimensions[ws.cell(start, i).column_letter].width = widths.get(k, 16)
    for c in ws[start]:
        c.font = Font(bold=True, color="FFFFFF")
        c.fill = PatternFill("solid", fgColor="1F4E79")
        c.alignment = Alignment(vertical="center", wrap_text=True)
    keys = [k for k, _ in cols]
    for row in ws.iter_rows(min_row=start + 1):
        for c in row:
            c.alignment = Alignment(vertical="top", wrap_text=True)
        if "result" in keys:
            rc = row[keys.index("result")]
            if rc.value in fills:
                rc.fill = PatternFill("solid", fgColor=fills[rc.value])
    ws.freeze_panes = ws.cell(start + 1, 1)
    if "verdict" in keys and rows:
        dv = DataValidation(type="list", formula1='"%s"' % ",".join(defaults.VERDICTS), allow_blank=True)
        col = ws.cell(start, keys.index("verdict") + 1).column_letter
        dv.add("%s%d:%s%d" % (col, start + 1, col, start + len(rows)))
        ws.add_data_validation(dv)

    wi = wb.create_sheet(INFO_SHEET)
    for k, v in info.items():
        wi.append([k, str(v)])
    wi.append([])
    wi.append(["판정 되받기", "「검토자 판정」 열에 채택/제외, 「보완 의견」 열에 기준 보완 의견을 적고 저장한 뒤 "
               "python review_tool.py feedback 이파일.xlsx 로 넣습니다. 「순번」 열과 이 시트는 지우지 마세요."])
    wi.column_dimensions["A"].width = 16
    wi.column_dimensions["B"].width = 90

    if prompts:
        wp = wb.create_sheet("판단 프롬프트")
        wp.append(["No", "프롬프트(셀을 복사해 사내 승인 AI 에 붙여 넣습니다)"])
        for no, text in prompts:
            wp.append([no, text[:32000]])
        wp.column_dimensions["A"].width = 8
        wp.column_dimensions["B"].width = 120
        for row in wp.iter_rows(min_row=2):
            for c in row:
                c.alignment = Alignment(vertical="top", wrap_text=True)
    wb.save(path)


def write_minutes_csv(path, rows, cfg, info):
    """openpyxl 이 없을 때의 대안 — 엑셀에서 열리는 UTF-8(BOM) CSV."""
    import csv
    cols = headers(cfg)
    with open(path, "w", newline="", encoding="utf-8-sig") as f:
        w = csv.writer(f)
        if info.get("예시 데이터") == "예":
            w.writerow(["예시 데이터 — 가상 표준 문서로 만든 시연용 결과입니다. 실제 심의 결과가 아닙니다."])
        w.writerow(["검토ID", info.get("검토ID", ""), "표준 파일", info.get("표준 파일", "")])
        w.writerow([h for _, h in cols])
        for r in rows:
            w.writerow([r.get(k, "") for k, _ in cols])


def parse_map(pairs):
    """["verdict=판정", "검토자 판정=판정"] → {key: 열 이름}. 키 또는 기본 열 이름 모두 받습니다."""
    by_header = {c["header"]: c["key"] for c in defaults.COLUMNS}
    out = {}
    for p in pairs or []:
        if "=" not in p:
            raise ConfigError("--map 은 「키=열이름」 형식입니다: %s" % p)
        k, v = [x.strip() for x in p.split("=", 1)]
        k = by_header.get(k, k)
        if k not in defaults.COLUMN_KEYS:
            raise ConfigError("--map 의 키 「%s」를 모릅니다. 쓸 수 있는 키: %s" % (k, ", ".join(defaults.COLUMN_KEYS)))
        out[k] = v
    return out


def locate_columns(header_row, cfg, overrides=None):
    """머리글 줄에서 필요한 열 위치를 찾습니다(열 매핑). 못 찾으면 무엇을 기대했는지 알려 줍니다."""
    want = {c["key"]: c["header"] for c in cfg.get("columns", [])}
    want.update(overrides or {})
    heads = [("" if h is None else str(h).strip()) for h in header_row]
    pos, missing = {}, []
    for key in ("seq", "no", "result", "detail", "verdict", "comment"):
        name = want.get(key)
        if name in heads:
            pos[key] = heads.index(name)
        elif key in ("seq", "verdict"):
            missing.append("%s(열 이름 「%s」)" % (key, name))
    if missing:
        raise ConfigError("회의록에서 필요한 열을 찾지 못했습니다: %s. 파일의 머리글: %s. "
                          "열 이름이 바뀌었으면 --map 키=열이름 으로 알려 주세요(예: --map verdict=판정)." % (", ".join(missing), " | ".join(h for h in heads if h)))
    return pos


def read_feedback(path, cfg, overrides=None):
    """(review_id, [{seq, no, result, detail, verdict, comment}], 경고 목록)."""
    openpyxl = _openpyxl()
    try:
        wb = openpyxl.load_workbook(path, data_only=True)
    except Exception as e:
        raise ConfigError("회의록 엑셀을 열 수 없습니다: %s (%s)" % (path, e))
    if INFO_SHEET not in wb.sheetnames:
        raise ConfigError("「정보」 시트가 없습니다. 이 도구가 만든 회의록 초안이 맞는지 확인해 주세요.")
    info = {}
    for row in wb[INFO_SHEET].iter_rows(values_only=True):
        if row and row[0]:
            info[str(row[0]).strip()] = "" if len(row) < 2 or row[1] is None else str(row[1]).strip()
    try:
        rid = int(info.get("검토ID", ""))
    except ValueError:
        raise ConfigError("「정보」 시트에 검토ID 가 없습니다.")
    ws = wb[MAIN_SHEET] if MAIN_SHEET in wb.sheetnames else wb.worksheets[0]
    rows = list(ws.iter_rows(values_only=True))
    hi = None
    last_err = None
    for i, r in enumerate(rows[:5]):
        try:
            pos = locate_columns(r, cfg, overrides)
            hi = i
            break
        except ConfigError as e:
            last_err = e
    if hi is None:
        raise last_err
    items, warnings = [], []
    for r in rows[hi + 1:]:
        def g(k):
            j = pos.get(k)
            return "" if j is None or j >= len(r) or r[j] is None else str(r[j]).strip()
        seq = g("seq")
        if not seq:
            continue
        try:
            seq = int(float(seq))
        except ValueError:
            warnings.append("순번 「%s」은(는) 숫자가 아니라 건너뜁니다." % seq)
            continue
        v = g("verdict")
        if v and v not in defaults.VERDICTS:
            warnings.append("순번 %d: 판정 「%s」은(는) 채택/제외가 아니라 비워 둡니다." % (seq, v))
            v = ""
        items.append({"seq": seq, "no": g("no"), "result": g("result"), "detail": g("detail"), "verdict": v, "comment": g("comment")})
    return rid, items, warnings
